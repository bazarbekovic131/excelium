"""Справочник подписантов: люди из Структуры, роли из должностей шлюза."""
import io
import json
import sqlite3
import tempfile
from pathlib import Path

import openpyxl

from conftest import docv_headers
from gateway.jobsqueue.db import connect, init_db
from gateway.signers import SignerStore, fio_key, print_name

MODEL = json.loads((Path(__file__).parent / "data" / "model.json").read_text(encoding="utf-8"))


def _render(client, kind, payload):
    r = client.post(f"/render/registry/{kind}", json=payload, headers=docv_headers())
    assert r.status_code == 200, r.text
    token = r.json()["download_url"].rsplit("/", 1)[1]
    return openpyxl.load_workbook(io.BytesIO(client.get(f"/files/{token}").content))


def _structura(client, items):
    r = client.post("/directory/structura", headers=docv_headers(), json={"items": items})
    assert r.status_code == 200


# --- первичное наполнение -------------------------------------------------

def test_seed_filled_from_old_sources(client):
    stats = client.app.state.signers.stats()
    assert stats["rules"] > 100 and stats["sets"] >= 5 and stats["roles"] > 20
    # одиннадцать гендиректоров — одиннадцать должностей с одним названием
    roles = client.app.state.signers.roles()
    gendirs = [r for r in roles if r["title"] == "Генеральный директор"]
    assert len(gendirs) == 11 and len({r["name"] for r in gendirs}) == 11
    # должность с одним человеком получает ключ без фамилии
    assert any(r["name"] == "Главный бухгалтер" for r in roles)


def test_seed_matches_vlookup_line_for_line(client):
    """Перенос обязан печатать то же, что печатал VLOOKUP по скрытому
    листу: сравниваем каждый лист тестового реестра."""
    import yaml
    from gateway.signers import normalize_name, normalize_object

    ws = openpyxl.load_workbook("templates/excel/template.xlsx")["СПР_ПОДПИСАНТОВ"]
    spr = {}
    for r in range(2, ws.max_row + 1):
        n = ws[f"B{r}"].value
        if isinstance(n, int):
            spr[n] = (str(ws[f"F{r}"].value or "").strip(), str(ws[f"H{r}"].value or "").strip(),
                      str(ws[f"J{r}"].value or "").strip())
    raw = yaml.safe_load(Path("data/signers_seed.yaml").read_text(encoding="utf-8"))
    rules = {}
    for rule in raw["rules"]:
        for obj in rule["objects"]:
            rules[(normalize_name(rule["company"]), normalize_object(obj))] = rule
    for fb in raw["company_fallbacks"]:
        rules.setdefault((normalize_name(fb["company"]), ""), fb)
    excl = raw["expense_type_exclusions"]

    def line(i):
        fio, comp, pos = spr[i]
        return (" ".join(x for x in (pos, comp) if x), fio)

    def vlookup(company, obj, zatraty):
        rule = rules.get((normalize_name(company), normalize_object(obj))) \
            or rules.get((normalize_name(company), ""))
        if rule is None:
            return [None, None], []
        ids = list(raw["approver_lists"][rule["approvers"]])
        if str(zatraty or "").strip().casefold() in excl["expense_types"]:
            ids = [i for i in ids if i not in excl["remove_ids"]]
        d = rule["directors"]
        return [line(n) if n and n in spr else None for n in (d[0], d[1])], \
            [line(i) for i in ids]

    wb = _render(client, "inner", MODEL)
    groups = {}
    for item in MODEL["request"]:
        groups.setdefault((item["organization"], item["object_name"]), []).append(item)
    underline = "_" * 28 + " "
    checked = 0
    for (company, obj), group in groups.items():
        exp_top, exp_lines = vlookup(company, obj, group[0].get("zatraty"))
        sheet = next(sh for sh in wb.worksheets if sh["G11"].value == obj
                     and company in str(sh["F17"].value))
        got_top = [(str(sheet["F3"].value or ""), str(sheet["F5"].value or "").replace(underline, "")),
                   (str(sheet["I3"].value or ""), str(sheet["I5"].value or "").replace(underline, ""))]
        got_top = [t if t[1] else None for t in got_top]
        got_lines = [(str(sheet.cell(row=rr, column=6).value), str(sheet.cell(row=rr, column=9).value))
                     for rr in range(18, sheet.max_row + 1)
                     if sheet.cell(row=rr, column=6).value and sheet.cell(row=rr, column=9).value
                     and not str(sheet.cell(row=rr, column=6).value).startswith("Заявитель: ")]
        assert got_top == exp_top, (company, obj)
        assert got_lines == exp_lines, (company, obj)
        checked += 1
    assert checked == len(groups)


# --- подбор ---------------------------------------------------------------

def test_resolve_endpoint_serves_docv(client):
    r = client.get("/signers", params={"company": "ТОО «Шар-Кұрылыс»",
                                       "object": "Администрация"}, headers=docv_headers())
    assert r.status_code == 200
    body = r.json()
    assert body["utverzhdayu"]["fio"] and body["utverzhdayu"]["position"]
    assert body["coordinators"] and body["source"] in ("объект", "компания")
    assert client.get("/signers").status_code == 403


def test_company_name_variations_still_find_signers(client):
    store = client.app.state.signers
    expected = store.resolve("ТОО «Шар-Кұрылыс»", "Администрация")
    assert expected["utverzhdayu"]
    for variant in ('ТОО «Шар-Кұрылыс» (KZT)', 'ТОО "Шар-Кұрылыс"', 'тоо «шар-курылыс»'):
        assert store.resolve(variant, "Администрация") == expected, variant
    unknown = store.resolve("ТОО «Никто»", "Нигде")
    assert unknown["coordinators"] == [] and unknown["utverzhdayu"] is None


def test_objects_with_brackets_stay_distinct(client):
    store = client.app.state.signers
    objects = {b["object_name"] for b in store.rules() if b["company"] == 'ТОО "СМУ Аргон"'}
    assert {"Школа (Нұра)", "Школа (Тельман)", "Школа (Уркер)"} <= objects


def test_object_without_own_top_inherits_company_rule(client):
    store = client.app.state.signers
    store.save_role("Директор Наследства", title="Директор", holder_name="Наследный Н.Н.")
    store.save_rule(company="ТОО «Наследство»", object_name="", set_name="list_1",
                    utverzhdayu="Директор Наследства")
    store.save_rule(company="ТОО «Наследство»", object_name="ЖК Первый", set_name="list_2")
    at_object = store.resolve("ТОО «Наследство»", "ЖК Первый")
    assert at_object["source"] == "объект"
    assert at_object["utverzhdayu"] == {"fio": "Наследный Н.Н.", "position": "Директор",
                                        "company": "ТОО «Наследство»"}
    # своё правило объекта сильнее
    store.save_role("Прораб объекта", holder_name="Свой С.С.")
    store.save_rule(company="ТОО «Наследство»", object_name="ЖК Первый", set_name="list_2",
                    utverzhdayu="Прораб объекта", utverzhdayu_company="ТОО «Своя»")
    own = store.resolve("ТОО «Наследство»", "ЖК Первый")["utverzhdayu"]
    assert own == {"fio": "Свой С.С.", "position": "Прораб объекта", "company": "ТОО «Своя»"}


def test_expense_type_skips_lines(client):
    store = client.app.state.signers
    with_all = store.resolve('ТОО "СМУ Аргон"', 'ЖК "New Line"', expense_type="СМР")
    salary = store.resolve('ТОО "СМУ Аргон"', 'ЖК "New Line"', expense_type="Зарплата")
    pto = "Начальник производственно-технического отдела"
    assert any(c["position"] == pto for c in with_all["coordinators"])
    assert not any(c["position"] == pto for c in salary["coordinators"])


# --- должности: один человек, ФИО из Структуры -----------------------------

def test_role_takes_fio_from_structura_and_title_from_itself(client):
    """Суть модели: человек — из Структуры, роль и печатаемое название —
    от должности шлюза. Сменился человек на должности — сменилась подпись
    во всех правилах, где на неё ссылаются."""
    store = client.app.state.signers
    _structura(client, [
        {"uid": "u-1", "display_name": "Аманов Бауыржан Шарипович (Генеральный директор)",
         "position": "Генеральный директор", "department": "Дирекция"},
        {"uid": "u-2", "display_name": "Второй Директор", "position": "Директор"}])
    store.save_role("Гендиректор Шар-Құрылыс", title="Генеральный директор", holder_uid="u-1")
    for company in ("ТОО «Одна»", "ТОО «Другая»"):
        store.save_rule(company=company, object_name="", set_name="list_1",
                        utverzhdayu="Гендиректор Шар-Құрылыс")
    first = store.resolve("ТОО «Одна»")["utverzhdayu"]
    # хвост «(Генеральный директор)» из display_name в подпись не идёт
    assert first == {"fio": "Аманов Бауыржан Шарипович", "position": "Генеральный директор",
                     "company": "ТОО «Одна»"}
    store.save_role("Гендиректор Шар-Құрылыс", title="Генеральный директор", holder_uid="u-2")
    assert store.resolve("ТОО «Одна»")["utverzhdayu"]["fio"] == "Второй Директор"
    assert store.resolve("ТОО «Другая»")["utverzhdayu"]["fio"] == "Второй Директор"


def test_same_title_different_people_are_different_roles(client):
    store = client.app.state.signers
    store.save_role("Гендиректор А", title="Генеральный директор", holder_name="А. А.")
    store.save_role("Гендиректор Б", title="Генеральный директор", holder_name="Б. Б.")
    store.save_rule(company="ТОО «А»", object_name="", set_name="list_1", utverzhdayu="Гендиректор А")
    store.save_rule(company="ТОО «Б»", object_name="", set_name="list_1", utverzhdayu="Гендиректор Б")
    a = store.resolve("ТОО «А»")["utverzhdayu"]
    b = store.resolve("ТОО «Б»")["utverzhdayu"]
    assert a["position"] == b["position"] == "Генеральный директор"
    assert a["fio"] == "А. А." and b["fio"] == "Б. Б."


def test_structura_name_wins_over_typed_name(client):
    store = client.app.state.signers
    _structura(client, [{"uid": "u-9", "display_name": "Из Структуры И.С."}])
    store.save_role("Проверяющий", holder_uid="u-9", holder_name="Вписанный В.В.")
    store.save_rule(company="ТОО «Кто»", object_name="", set_name="list_1",
                    soglasovano="Проверяющий")
    assert store.resolve("ТОО «Кто»")["soglasovano"]["fio"] == "Из Структуры И.С."
    # uid исчез из Структуры — остаётся вписанное имя, а не пустота
    _structura(client, [{"uid": "u-other", "display_name": "Другой Д.Д."}])
    assert store.resolve("ТОО «Кто»")["soglasovano"]["fio"] == "Вписанный В.В."


def test_vacant_role_signs_nobody(client):
    store = client.app.state.signers
    store.save_role("Пустая должность", title="Никто")
    store.save_rule(company="ТОО «Пусто»", object_name="", set_name="list_1",
                    utverzhdayu="Пустая должность")
    assert store.resolve("ТОО «Пусто»")["utverzhdayu"] is None
    assert any(r["name"] == "Пустая должность" and r["vacant"] for r in store.roles())


def test_role_rename_follows_references_and_delete_is_guarded(client):
    store = client.app.state.signers
    store.save_role("Старое имя", holder_name="Кто-то К.К.")
    store.save_rule(company="ТОО «Переименование»", object_name="", set_name="list_1",
                    utverzhdayu="Старое имя")
    store.rename_role("Старое имя", "Новое имя")
    assert store.resolve("ТОО «Переименование»")["utverzhdayu"]["position"] == "Новое имя"
    assert store.role_usage("Новое имя") == 1
    try:
        store.delete_role("Новое имя")
        assert False, "используемую должность удалять нельзя"
    except ValueError:
        pass
    store.save_rule(company="ТОО «Переименование»", object_name="", set_name="list_1")
    store.delete_role("Новое имя")
    assert not any(r["name"] == "Новое имя" for r in store.roles())


def test_link_by_name_fills_uid_once(client):
    store = client.app.state.signers
    assert store.resolve("ТОО «Шар-Кұрылыс»", "Администрация")["utverzhdayu"]["fio"] \
        == "Аманов Б.Ш."
    _structura(client, [{"uid": "u-1", "display_name": "Аманов Бауыржан Шарипович",
                         "position": "Генеральный директор"}])
    result = store.link_by_name()
    assert result["linked"] >= 1
    after = store.resolve("ТОО «Шар-Кұрылыс»", "Администрация")["utverzhdayu"]
    assert after["fio"] == "Аманов Бауыржан Шарипович"
    assert after["position"] == "Генеральный директор"


def test_fio_helpers():
    assert fio_key("Аманов Б.Ш.") == fio_key("Аманов Бауыржан Шарипович")
    assert fio_key("Аманов Б.Ш.") == fio_key("Аманов Бауыржан Шарипович (Гендиректор)")
    assert fio_key("Аманов Б.Ш.") != fio_key("Аманов Дархан Ерланович")
    assert print_name("Иванов И.И. (Прораб)") == "Иванов И.И."
    assert print_name("  ДРС ") == "ДРС"


# --- наборы, реестр предстоящих платежей -----------------------------------

def test_priority_registry_takes_coordinators_from_store(client):
    entry = dict(MODEL["request"][0], organization="ТОО «Шар-Кұрылыс»",
                 object_name="Администрация", payment_number="1", status="в работе")
    wb = _render(client, "priority", {"request": [entry]})
    sheet = wb[wb.sheetnames[0]]
    lines = [str(sheet.cell(row=r, column=2).value)
             for r in range(9, sheet.max_row + 1) if sheet.cell(row=r, column=2).value]
    expected = client.app.state.signers.resolve("ТОО «Шар-Кұрылыс»", "Администрация")
    assert lines and len(lines) == len(expected["coordinators"])
    assert lines[0] == expected["coordinators"][0]["position"]


def test_set_delete_guarded(client):
    store = client.app.state.signers
    try:
        store.delete_set("list_1")
        assert False
    except ValueError:
        pass
    store.save_set("временный", [{"role": "Главный бухгалтер"}])
    assert "временный" in store.sets()
    store.delete_set("временный")
    assert "временный" not in store.sets()


# --- перенос с прежней модели ---------------------------------------------

def test_legacy_tables_migrate_into_roles(tmp_path):
    """База прошлой недели: люди отдельной таблицей и четыре способа
    указать подписанта. После переноса — только должности, и реестр
    печатается так же."""
    db = tmp_path / "gateway.db"
    init_db(db)
    with connect(db) as conn:
        conn.executescript("""
        CREATE TABLE signer_people (id INTEGER PRIMARY KEY, fio TEXT, fio_key TEXT,
            position TEXT, docv_uid TEXT, updated_at TEXT);
        CREATE TABLE signer_positions (name TEXT PRIMARY KEY, holder_uid TEXT,
            holder_person_id INTEGER, updated_at TEXT);
        CREATE TABLE signer_sets (name TEXT, ord INTEGER, person_id INTEGER, position TEXT,
            position_ref TEXT, dept_ref TEXT, print_company TEXT, mark TEXT,
            skip_expense_types TEXT);
        CREATE TABLE signer_bindings (id INTEGER PRIMARY KEY, company TEXT, company_key TEXT,
            object_name TEXT, object_key TEXT, set_name TEXT,
            soglasovano_id INTEGER, soglasovano_position TEXT, soglasovano_company TEXT,
            soglasovano_ref TEXT, soglasovano_dept TEXT,
            utverzhdayu_id INTEGER, utverzhdayu_position TEXT, utverzhdayu_company TEXT,
            utverzhdayu_ref TEXT, utverzhdayu_dept TEXT);
        INSERT INTO signer_people VALUES (1,'Аманов Б.Ш.','аманов бш','Генеральный директор','',''),
                                         (2,'Омарова Г.А.','омарова га','Главный бухгалтер','','');
        INSERT INTO signer_positions VALUES ('Финансовый директор','u-fin',NULL,'');
        INSERT INTO signer_sets VALUES ('list_1',0,2,'','','','','',''),
                                       ('list_1',1,NULL,'','gw:Финансовый директор','','','',''),
                                       ('list_1',2,NULL,'','u-str','','ТОО «Стр»','СОГЛАСОВАНО','');
        INSERT INTO signer_bindings VALUES (1,'ТОО «Тест»','тоо тест','','','list_1',
            NULL,'','','','', 1,'Генеральный директор','ТОО "Шар Құрылыс"','','');
        INSERT INTO directories VALUES ('structura','u-fin','{"name":"Финансист Ф.Ф.","position":"Финансовый директор"}',''),
                                       ('structura','u-str','{"name":"Структурный С.С.","position":"Юрист"}','');
        """)
    store = SignerStore(db)
    moved = store.migrate_legacy()
    assert moved == 1
    with connect(db) as conn:
        left = {r["name"] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert not {"signer_people", "signer_sets", "signer_bindings", "signer_positions"} & left

    resolved = store.resolve("ТОО «Тест»")
    assert resolved["utverzhdayu"] == {"fio": "Аманов Б.Ш.", "position": "Генеральный директор",
                                       "company": 'ТОО "Шар Құрылыс"'}
    assert [c["fio"] for c in resolved["coordinators"]] == \
        ["Омарова Г.А.", "Финансист Ф.Ф.", "Структурный С.С."]
    assert resolved["coordinators"][2]["mark"] == "СОГЛАСОВАНО"
    assert resolved["coordinators"][2]["position"] == "Юрист"
    roles = {r["name"]: r for r in store.roles()}
    assert roles["Финансовый директор"]["holder_uid"] == "u-fin"
    assert roles["Юрист"]["holder_uid"] == "u-str"
    # повторный запуск ничего не ломает
    assert store.migrate_legacy() == 0


# --- выключатели ------------------------------------------------------------

def test_disabled_role_falls_back_to_company_rule(client):
    store = client.app.state.signers
    store.save_role("Директор компании", title="Директор", holder_name="Компанейский К.К.")
    store.save_role("Директор объекта", title="Директор", holder_name="Объектный О.О.")
    store.save_rule(company="ТОО «Отпуск»", object_name="", set_name="list_1",
                    utverzhdayu="Директор компании")
    store.save_rule(company="ТОО «Отпуск»", object_name="Стройка", set_name="list_1",
                    utverzhdayu="Директор объекта")
    assert store.resolve("ТОО «Отпуск»", "Стройка")["utverzhdayu"]["fio"] == "Объектный О.О."
    store.set_roles_enabled(["Директор объекта"], False)
    assert store.resolve("ТОО «Отпуск»", "Стройка")["utverzhdayu"]["fio"] == "Компанейский К.К."
    store.set_roles_enabled(["Директор компании"], False)
    assert store.resolve("ТОО «Отпуск»", "Стройка")["utverzhdayu"] is None
    store.set_roles_enabled(["Директор объекта", "Директор компании"], True)
    assert store.resolve("ТОО «Отпуск»", "Стройка")["utverzhdayu"]["fio"] == "Объектный О.О."


def test_disabled_set_line_is_skipped(client):
    store = client.app.state.signers
    lines = [{"role": e["role"], "company": e["print_company"], "mark": e["mark"],
              "skip_expense_types": e["skip_expense_types"], "enabled": "1"}
             for e in store.sets()["list_1"]]
    before = store.resolve("ТОО «Шар-Кұрылыс»", "Администрация")["coordinators"]
    lines[0]["enabled"] = "0"
    store.save_set("list_1", lines)
    after = store.resolve("ТОО «Шар-Кұрылыс»", "Администрация")["coordinators"]
    assert len(after) == len(before) - 1 and after[0] == before[1]
    assert store.sets()["list_1"][0]["enabled"] == 0   # строка на месте, но выключена


def test_disabled_rules(client):
    store = client.app.state.signers
    store.save_role("Ктото", holder_name="Кто-то К.К.")
    company_id = store.save_rule(company="ТОО «Правила»", object_name="", set_name="list_1",
                                 utverzhdayu="Ктото")
    object_id = store.save_rule(company="ТОО «Правила»", object_name="Объект",
                                set_name="list_2", utverzhdayu="Ктото")
    assert store.resolve("ТОО «Правила»", "Объект")["source"] == "объект"
    store.set_rules_enabled([object_id], False)
    assert store.resolve("ТОО «Правила»", "Объект")["source"] == "компания"
    store.set_rules_enabled([company_id], False)
    assert store.resolve("ТОО «Правила»", "Объект")["source"] == "нет правила"
    # отключённое правило видно в списке, но не действует
    assert not next(r for r in store.rules() if r["id"] == object_id)["enabled"]


def test_roles_save_keeps_enabled_flag(client):
    store = client.app.state.signers
    store.save_role("Отпускник", holder_name="Отпускник О.О.")
    store.set_roles_enabled(["Отпускник"], False)
    store.save_role("Отпускник", title="Новое название", holder_name="Отпускник О.О.")
    assert not next(r for r in store.roles() if r["name"] == "Отпускник")["enabled"]
    store.save_role("Отпускник", holder_name="Отпускник О.О.", enabled=True)
    assert next(r for r in store.roles() if r["name"] == "Отпускник")["enabled"]


def test_apply_to_selected_rules_only(client):
    store = client.app.state.signers
    rules = store.company_rules('ТОО "СМУ Аргон"')
    store.save_role("Избранный", holder_name="Избранный И.И.")
    chosen = [r["id"] for r in rules[:2]]
    changed = store.apply_to_company('ТОО "СМУ Аргон"', utverzhdayu={"role": "Избранный"},
                                     rule_ids=chosen)
    assert changed == 2
    after = {r["id"]: r["utverzhdayu_role"] for r in store.company_rules('ТОО "СМУ Аргон"')}
    assert all(after[i] == "Избранный" for i in chosen)
    assert any(v != "Избранный" for i, v in after.items() if i not in chosen)


def test_delete_roles_skips_used(client):
    store = client.app.state.signers
    used = next(r["name"] for r in store.roles() if r["used"])
    store.save_role("Свободная", holder_name="Никто Н.Н.")
    deleted, skipped = store.delete_roles([used, "Свободная"])
    assert deleted == 1 and skipped == [used]


# --- экспорт и импорт -------------------------------------------------------

def test_export_import_roundtrip(client):
    store = client.app.state.signers
    store.set_roles_enabled([store.roles()[0]["name"]], False)
    before = store.resolve("ТОО «Шар-Кұрылыс»", "Администрация")
    data = store.export_json()
    assert data["version"] == 1 and data["roles"] and data["sets"] and data["rules"]
    assert any(r["enabled"] is False for r in data["roles"])
    # стираем всё и возвращаем из файла
    with connect(client.settings.db_path) as conn:
        for table in ("signer_rules", "signer_set_lines", "signer_roles"):
            conn.execute(f"DELETE FROM {table}")
    assert store.resolve("ТОО «Шар-Кұрылыс»", "Администрация")["source"] == "нет правила"
    counts = store.import_json(json.loads(json.dumps(data)))
    assert counts["rules"] == len(data["rules"]) and counts["roles"] == len(data["roles"])
    assert store.resolve("ТОО «Шар-Кұрылыс»", "Администрация") == before
    assert not store.roles()[0]["enabled"]
    for bad in ({"version": 2}, {"version": 1, "roles": {}, "sets": {}, "rules": []}, []):
        try:
            store.import_json(bad)
            assert False, bad
        except ValueError:
            pass
